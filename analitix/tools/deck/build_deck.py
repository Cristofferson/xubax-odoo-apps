# -*- coding: utf-8 -*-
"""Genera el PPT de presentación de ANALITIX.

Criterios que pidió el usuario: muy vendedora, poco texto, más imagen y
gráfica. De ahí las decisiones:

* Una idea por lámina y como máximo una frase. Si algo necesita explicación,
  la da quien presenta — la lámina no es el guion.
* Las imágenes son los renders en español SIN el titular quemado, para que el
  título de la lámina no compita con texto dentro de la foto.
* Las cifras son las REALES de la demo de retail. Si el cliente pide ver el
  sistema después, la lámina y la pantalla tienen que decir lo mismo.
* Sale un .pptx de verdad, editable y presentable sin internet: se va a abrir
  en la computadora de un cliente, quizá en una tienda con mala señal.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TINTA = RGBColor(0x12, 0x15, 0x1C)
ORO = RGBColor(0xC9, 0xA2, 0x27)
HUESO = RGBColor(0xEE, 0xF0, 0xF4)
GRIS = RGBColor(0x9A, 0xA3, 0xB4)
MENTA = RGBColor(0x1A, 0xD3, 0xBB)
CIRUELA = RGBColor(0x98, 0x51, 0x84)

IMG = "/tmp/ppt/"
ANCHO, ALTO = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = ANCHO, ALTO
BLANCO = prs.slide_layouts[6]          # en blanco


def lamina(fondo=TINTA):
    s = prs.slides.add_slide(BLANCO)
    f = s.shapes.add_shape(1, 0, 0, ANCHO, ALTO)   # 1 = rectángulo
    f.fill.solid(); f.fill.fore_color.rgb = fondo
    f.line.fill.background()
    f.shadow.inherit = False
    return s


def texto(s, x, y, w, h, contenido, tam=28, color=HUESO, negrita=False,
          alineado=PP_ALIGN.LEFT, interlinea=1.15):
    caja = s.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, linea in enumerate(contenido.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = linea
        p.alignment = alineado
        p.line_spacing = interlinea
        for run in p.runs:
            run.font.size = Pt(tam)
            run.font.bold = negrita
            run.font.color.rgb = color
            run.font.name = "Arial"
    return caja


def foto_completa(s, archivo):
    """La imagen ocupa toda la lámina. Los renders son 16:9 igual que la
    diapositiva, así que no hay recorte ni deformación."""
    s.shapes.add_picture(IMG + archivo, 0, 0, width=ANCHO, height=ALTO)


def banda(s, alto_pulgadas=2.35, arriba=False):
    """Franja oscura sobre la foto para que el texto se lea sin pelear con
    ella. Sin esto, media frase cae sobre una zona clara y desaparece."""
    h = Inches(alto_pulgadas)
    y = 0 if arriba else ALTO - h
    r = s.shapes.add_shape(1, 0, y, ANCHO, h)
    r.fill.solid(); r.fill.fore_color.rgb = TINTA
    r.line.fill.background(); r.shadow.inherit = False
    return r


def kicker(s, txt, x=Inches(0.85), y=Inches(0.7)):
    texto(s, x, y, Inches(9), Inches(0.5), txt.upper(), tam=16, color=ORO,
          negrita=True)


# ─────────────────────────────────────────────────────────── 1 · portada
s = lamina()
s.shapes.add_picture(IMG + "analitix-logo-blanco.png", Inches(0.9), Inches(2.5),
                     width=Inches(4.6))
texto(s, Inches(0.95), Inches(4.0), Inches(9.5), Inches(2),
      "Lo que tu tienda física\nnunca te ha dicho.", tam=48, negrita=True)
texto(s, Inches(1.0), Inches(6.3), Inches(8), Inches(0.6),
      "XUBAX · Tecnología in situ", tam=18, color=GRIS)

# ─────────────────────────────────────────────── 2 · la pregunta incómoda
s = lamina()
kicker(s, "empecemos por aquí")
# Una sola caja con dos párrafos: en cajas separadas, si la primera envuelve
# a dos líneas se le encima a la segunda.
caja = texto(s, Inches(0.85), Inches(2.1), Inches(11.6), Inches(2.6),
             "Vendiste $80,000 el sábado.", tam=50, negrita=True)
p2 = caja.text_frame.add_paragraph()
p2.text = "¿Estuvo bien o estuvo mal?"
for r in p2.runs:
    r.font.size = Pt(50); r.font.bold = True; r.font.color.rgb = ORO
    r.font.name = "Arial"
texto(s, Inches(0.9), Inches(5.3), Inches(10), Inches(1.5),
      "Sin saber cuánta gente entró, esa cifra no significa nada.",
      tam=24, color=GRIS)

# ─────────────────────────────────────────── 3 · el mismo número, dos negocios
s = lamina()
kicker(s, "la misma venta")
for i, (n, txt, col) in enumerate([
        ("40", "entraron\n30 compraron", MENTA),
        ("400", "entraron\n30 compraron", RGBColor(0xF0, 0x71, 0x3F))]):
    x = Inches(1.0 + i * 6.0)
    texto(s, x, Inches(2.1), Inches(5), Inches(2), n, tam=120, negrita=True, color=col)
    texto(s, x, Inches(4.4), Inches(5), Inches(1.6), txt, tam=26)
texto(s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(1),
      "Dos negocios completamente distintos. Hoy los dos se ven igual en tu ERP.",
      tam=24, color=GRIS)

# ────────────────────────────────────────── 4-7 · la historia, minuto a minuto
for archivo, hora, frase in [
        ("door.jpg", "14:02", "Entra."),
        ("floor.jpg", "14:04", "Se detiene en el mostrador de anillos."),
        ("frustrated.jpg", "14:06", "Cuatro minutos. Nadie se ha acercado."),
        ("walkout.jpg", "14:11", "Se va. Sin nada.")]:
    s = lamina()
    foto_completa(s, archivo)
    banda(s)
    texto(s, Inches(0.85), Inches(5.5), Inches(2.2), Inches(1),
          hora, tam=40, negrita=True, color=ORO)
    texto(s, Inches(3.2), Inches(5.55), Inches(9.3), Inches(1.5),
          frase, tam=34, negrita=True)

# ───────────────────────────────────────────────────── 8 · el número del mes
s = lamina()
texto(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(2.4),
      "62", tam=200, negrita=True, color=ORO)
texto(s, Inches(1.05), Inches(4.75), Inches(11.4), Inches(2),
      "clientes se fueron sin comprar en 14 días.\nNinguno dejó rastro en un reporte.",
      tam=30, negrita=True)

# ────────────────────────────────────────────────────── 9 · lo que se rescata
s = lamina()
s.shapes.add_picture(IMG + "chart-salidas.png", 0, Inches(0.33), width=ANCHO)

# ───────────────────────────────────────────────────────── 10 · el aviso
s = lamina()
kicker(s, "con analitix, a las 14:06")
texto(s, Inches(0.85), Inches(1.25), Inches(5.7), Inches(1.2),
      "Vibra un celular.", tam=42, negrita=True)
texto(s, Inches(0.9), Inches(2.6), Inches(5.6), Inches(3),
      "El vendedor que cubre esa zona recibe el aviso mientras la clienta "
      "sigue en la tienda.\n\nElla nunca se entera.", tam=24, color=GRIS)
s.shapes.add_picture(IMG + "alert-discuss.jpg", Inches(6.9), Inches(1.3),
                     width=Inches(5.7))

# ─────────────────────────────────────────────────────── 11 · la gráfica
s = lamina()
s.shapes.add_picture(IMG + "chart-hora.png", 0, Inches(0.33), width=ANCHO)

# ───────────────────────────────────────────────────── 12 · qué se instala
s = lamina()
kicker(s, "qué se instala en la tienda")
texto(s, Inches(0.85), Inches(1.18), Inches(11.6), Inches(0.9),
      "Una mini PC y tu cámara de siempre.", tam=34, negrita=True)
s.shapes.add_picture(IMG + "install.jpg", Inches(0.85), Inches(2.5), width=Inches(5.7))
s.shapes.add_picture(IMG + "cameras.jpg", Inches(6.85), Inches(2.5), width=Inches(5.7))
texto(s, Inches(0.85), Inches(5.95), Inches(5.7), Inches(1.2),
      "Va en la trastienda. Sin tarjeta de video, nada en la caja, nada en la nube.",
      tam=17, color=GRIS)
texto(s, Inches(6.85), Inches(5.95), Inches(5.7), Inches(1.2),
      "Análoga, webcam o IP: las tres sirven. No hay que recablear nada.",
      tam=17, color=GRIS)

# ──────────────────────────────────────────────────────── 13 · privacidad
s = lamina()
kicker(s, "antes de que lo preguntes")
texto(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(1.4),
      "Ninguna imagen sale de tu tienda.", tam=38, negrita=True, color=MENTA)
texto(s, Inches(0.9), Inches(3.5), Inches(11), Inches(2.0),
      "El reconocimiento ocurre en esa computadora y el cuadro de video se "
      "destruye ahí mismo.\nLo que viaja a Odoo son números.", tam=28, color=GRIS)
texto(s, Inches(0.9), Inches(5.7), Inches(11), Inches(1.2),
      "No graba · No identifica a nadie · No vigila a tu equipo",
      tam=22, negrita=True)

# ─────────────────────────────────────────────────────────── 14 · niveles
s = lamina()
kicker(s, "cómo se contrata")
texto(s, Inches(0.85), Inches(1.15), Inches(11.5), Inches(1),
      "Tres niveles, por sucursal", tam=40, negrita=True)
niveles = [("CONTEO", "Cuánta gente entra\ny tu conversión real", CIRUELA),
           ("ANALÍTICOS VISUALES", "Todo lo anterior, más\nzonas, visitas y exhibidores", MENTA),
           ("ACCIONES", "Todo lo anterior, más\nel aviso que rescata la venta", ORO)]
for i, (nombre, det, col) in enumerate(niveles):
    x = Inches(0.85 + i * 4.05)
    caja = s.shapes.add_shape(1, x, Inches(2.5), Inches(3.75), Inches(2.9))
    caja.fill.solid(); caja.fill.fore_color.rgb = RGBColor(0x1B, 0x20, 0x2B)
    caja.line.fill.background(); caja.shadow.inherit = False
    texto(s, x + Inches(0.3), Inches(2.8), Inches(3.2), Inches(0.8),
          nombre, tam=19, negrita=True, color=col)
    texto(s, x + Inches(0.3), Inches(3.7), Inches(3.2), Inches(1.6), det, tam=18)
texto(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(1),
      "Cada nivel contiene al anterior: no se suman. Y el nivel es por tienda, "
      "así que tu sucursal insignia puede ir en Acciones y las chicas en Conteo.",
      tam=19, color=GRIS)

# ──────────────────────────────────────────────────────────── 15 · cierre
s = lamina()
texto(s, Inches(0.85), Inches(2.4), Inches(11.6), Inches(2),
      "Veámoslo con los números\nde tu tienda.", tam=52, negrita=True,
      interlinea=1.05)
texto(s, Inches(0.9), Inches(4.7), Inches(10), Inches(1.5),
      "Una demostración con datos reales de ejemplo. De ahí sale el alcance.",
      tam=26, color=GRIS)
texto(s, Inches(0.9), Inches(6.2), Inches(11), Inches(0.8),
      "XUBAX · xubax.com/analitix · 443 688 7346", tam=20, color=ORO, negrita=True)

prs.save("/tmp/ppt/ANALITIX.pptx")
print("láminas:", len(prs.slides.__iter__.__self__._sldIdLst))
